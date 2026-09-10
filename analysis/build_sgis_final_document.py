"""Build the submission from the maintained draft, including numbered evidence."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_sgis_submission_docx import configure_document, add_heading, add_paragraph, add_table, set_east_asia_font, FONT_BODY, FONT_HEADING

ROOT = Path(__file__).resolve().parents[1]
DRAFT = ROOT / 'docs/submission/sgis-uiseong-final-draft.md'
OUTPUT = ROOT / 'docs/submission/sgis-uiseong-revised-review.docx'

def clean(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')

def render(doc, text, appendix=False):
    lines = text.splitlines()
    i = 0
    code = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or line.startswith('<!--'): continue
        if line.startswith('```'):
            code = not code
            continue
        if line.startswith('## '):
            add_heading(doc, line[3:])
        elif line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_east_asia_font(p.add_run(line[2:]), FONT_HEADING, 16, bold=True)
        elif line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].startswith('|'):
                rows.append(lines[i]); i += 1
            cells = [[clean(c.strip()) for c in row.strip('|').split('|')] for row in rows if not re.fullmatch(r'[|:\-\s]+', row)]
            n = len(cells[0]); widths = [9638//n]*n; widths[-1] += 9638-sum(widths)
            if n == 3: widths = [2300, 1800, 5538]
            add_table(doc, cells[0], cells[1:], widths)
        elif line.startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic = p.add_run().add_picture(str(DRAFT.parent / match[2]), width=Inches(6.2))
            pic._inline.docPr.set('descr', match[1])
        else:
            is_caption = bool(re.match(r'\[(표|그림) \d+\.', line))
            is_source = line.startswith('출처:') or line.startswith('- [출처 ')
            p = add_paragraph(doc, clean(line), size=11 if appendix or is_caption or is_source or code else 15, after=4)
            if appendix or is_caption or is_source or code: p.paragraph_format.line_spacing = 1.2
            if is_caption:
                p.paragraph_format.keep_with_next = True
                for run in p.runs: run.bold = True

def build():
    text = DRAFT.read_text(encoding='utf-8')
    body, appendix = text.split('<!-- 제출 본문 끝.', 1)
    appendix = appendix.split('-->', 1)[1]
    # The editable evidence draft remains complete; this is the concise five-page form.
    pages = [
'''# SGIS 생활권역 통계를 활용한 의성 전통 못 관리와 청년 참여 정책 제안

## □ 추진배경 및 필요성
의성군 인구는 2024년 12월 48,690명에서 2025년 12월 47,902명으로 788명 감소했다. 지역자산을 관리할 인력과 활용 기반을 살피려면 시설 위치에 인구구조와 이동권역의 정보를 더해야 한다. [출처 2]

의성 전통수리 농업시스템은 2018년 국가중요농업유산 제10호로 지정되었다. 농림축산식품부·해양수산부는 2025.11.2. 공식 발표에서 FAO 세계중요농업유산(GIAHS) 등재 추진 대상으로 의성을 명시했다. [출처 3] [출처 5]

본 연구는 2025년 의성청년연구자 활동을 계기로 전통 못 427개를 지역자원으로 검토했다. 핵심 질문은 “각 못의 주변 여건을 근거로 어떤 조사와 정책 협의를 먼저 시작할 것인가”이다. 개별 못 모두를 유휴시설 또는 지정 농업유산으로 전제하지 않았다.

## □ 추진과정 및 활용분야
427개 못의 ID·좌표를 정리하고 SGIS 2024년 기준 5분·10분 주행생활권 인구를 연결했다. 행정안전부 2024·2025년 12월 읍·면 인구구조와 VWorld 도로·주변시설을 결합했다. 공공정책 및 연구·교육 분야에서 활용할 시설별 근거표와 후속 조사 양식을 작성하였다.

청년은 만 19~39세, 고령은 만 65세 이상이다. 읍·면 비율은 지역 배경, SGIS 인구는 각 못의 이동권역 통계로 해석하였다. [출처 1] [출처 2] [출처 6]''',
'''## □ 추진내용
1. SGIS를 이용한 시설별 생활권 비교

427개 중 대표점·인근·중복좌표 47개를 비교에서 제외했다. 좌표 조건을 통과한 380개 중 5분 인구는 175개, 10분 인구는 190개, 양쪽 인구는 175개에서 확보했다. 값이 없는 시설을 0명으로 처리하지 않았다.

[표 1. SGIS 활용 전후의 검토 근거]

| 검토 질문 | 활용 전 | 활용 후 |
|---|---|---|
| 생활권 인구 규모 | 읍·면 인구와 공간조건 중심 | 190개 못에 10분 인구 연결 |
| 이동 범위의 차이 | 시설별 인구 비교 근거 없음 | 175개에서 5분→10분 비교 |
| 추가 조사 대상 | 위치와 주변환경 중심 | 좌표 확인 47개·10분 인구 미확보 190개 분리 |

출처: SGIS, 2024년 기준 저장 응답 [출처 1]; 작성자 분석.

10분 인구 제3사분위수 이상은 49개, 5분→10분 증가폭 제3사분위수 이상은 44개이며 중복은 42개다. 기준은 각각 190개와 175개에서 산출했다. 생활권이 겹치므로 인구를 합산하지 않았으며, 인구 규모를 실제 방문객 수로 간주하지 않았다.

산점도의 ID 474(단북면 노연리)는 5분 2,298명·10분 3,877명으로 저장 응답과 일치한다. 단일 좌표이고 입력 오류 근거가 없어 유지했다. X축 로그 변환으로 밀집 구간을 구분하고 최댓값을 주석 처리했다. [그림 2]''',
'''## □ 추진내용
2. 지역 여건을 결합한 정책 검토

SGIS 인구와 읍·면 청년·고령 비율, 주변시설, 좌표 품질을 연결했다. 아래 순서대로 판정하여 각 못은 하나의 정책 검토 그룹에 포함된다.

[표 2. 정책 검토 맥락과 후속 사업 협의]

| 구분 | 못 수 | 근거와 후속 검토 |
|---|---|---|
| 현장 확인 필요 | 47 | 좌표 조건 미충족; 위치·관리주체 확인 |
| 10분 인구 미확보 | 190 | 인구값 재수집; 낮은 수요로 해석하지 않음 |
| 공동체 지원 검토 | 53 | 고령 상위·주변 매핑시설 없음; 주민 관리수요 확인 |
| 청년참여 검토 | 43 | 고령 상위 외 청년비율 상위; 조사·기록·콘텐츠 활동 협의 |
| 접근·도달 검토 | 20 | 앞선 그룹 외 인구·증가폭 상위; 이동·체험 여건 확인 |
| 맥락 참고 | 74 | 추가 자료 확보 후 검토 |

출처: SGIS·행정안전부 및 본 연구 공간분석 [출처 1] [출처 2] [출처 6]; 작성자 분석.

‘상위’는 대상 못에 연결된 값의 제3사분위수 기준이다. 주변시설 0개는 확보한 레이어의 객체가 없다는 뜻이다. 최근접 도로거리는 UQ151 객체까지의 직선거리이며 주행시간과 구분했다. 보조 규칙기반 활용유형은 SGIS 인구를 판정조건에 사용하지 않아 별첨에 분리하였다.''',
'''## □ 추진성과
시설별 인구 비교 근거와 정책 협의 자료를 만들었다

SGIS를 통해 읍·면 총계로는 구분하기 어려운 개별 못의 주행생활권 인구와 확장 규모를 비교했다. 접근·도달 20개, 청년참여 43개, 공동체 지원 53개를 합한 116개에 관리상태·사진·조사일·의견을 입력할 현장검토 양식을 작성했다. [출처 6]

분석 결과를 지역자원 관리와 청년 참여 사업으로 연결했다

‘청년 못 조사단’·‘못 기록화 사업’을 통해 위치·관리상태 조사, 디지털 지도·기록, 주민 구술과 세대 간 지식 전승을 연계하는 방안을 제안했다. 생태 모니터링·스토리텔링·체험 프로그램으로 확장하고, 유지·보수와 복원은 관리주체 및 안전·건설 담당의 전문 검토에 연결하도록 구성했다.

관내 청년에게는 현장 경험과 농업·생태·지역문화 학습, 관외 청년에게는 전통문화·생태 체험과 공동체 참여 기회를 제공하는 것을 목표로 한다. 청년정책·안전건설·환경·농촌활력·문화관광 기능의 협업 사업으로 발전시킬 수 있다.

공식 실행 사례로 사업 방향의 가능성을 보강했다

농식품부는 2026.7.9. 의성 금성면에서 대학생·교직원 약 40명이 게임·웹소설·SNS 숏폼 제작과 마을환경 정비 등에 참여한다고 발표했다. 대학의 2026.7.20. 결과 게시문도 이를 뒷받침한다. 본 연구와 별도로 실행된 유사 사례이며, 본 연구의 사업 집행 실적으로 합산하지 않았다. [출처 3] [출처 4]''',
'''## □ 기대효과 및 확산방향
다른 기관이 재사용할 분석·검토 절차

자원 목록 구축 → 좌표·자료 품질 점검 → SGIS 생활권역 통계 결합 → 지역별 검토 기준 설정 → 현장 확인의 절차를 제안한다. 재사용할 요소는 시설 ID별 자료 결합, 응답·제외 사유 구분, 검토 근거표와 출처 기록이다. 각 지역의 자원 목록·경계·기준연도·이동시간과 정책 조건은 다시 설정해야 한다.

소규모 목록으로 SGIS 응답과 좌표 품질을 먼저 점검하고, 담당자와 판정 근거를 검토한 뒤 범위를 넓힌다. 기관 채택, 타 지역 독립 재현, 교육 실시는 후속 검증 과제다. 청년 정착·고용·조사비용 절감은 참여 지속률, 조사 기록 완성도, 주민 평가 등을 통해 사업 시행 후 측정한다.

SGIS 활용의 차별적 요소

전통 못을 고정된 시설 목록으로 보는 데서 나아가, 개별 자원의 이동권역 인구를 청년참여·공동체 관리의 검토 질문에 연결했다. 인구 규모뿐 아니라 근거가 부족한 이유까지 같은 ID로 남겨 후속 조사의 출발점을 제시했다.

심사기준에 대응하는 성과 범위

충실성(30점)은 SGIS 수집·비교·정책 근거 연결, 효과성(30점)은 연구자료와 116개 현장검토 양식, 확산가능성(30점)은 재사용 절차와 지역별 변경조건, 창의성(10점)은 전통자원·생활권 인구·청년 참여를 결합한 활용 방식으로 제시한다. [출처 7]

이하 별첨에는 실제 분석표·그림·원자료 대조·번호별 출처를 수록하였다.'''
    ]
    doc = Document()
    configure_document(doc)
    for i, page in enumerate(pages):
        if i: doc.add_page_break()
        render(doc, page)
    # Evidence is outside the five-page body. Remove internal editorial instructions.
    appendix = re.sub(r'## 편집 메모:.*?(?=## 별첨 4\.)', '', appendix, flags=re.S)
    for chunk in re.split(r'(?=^## 별첨 \d+\.)', appendix, flags=re.M):
        if not chunk.strip(): continue
        doc.add_page_break()
        render(doc, chunk, appendix=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))

if __name__ == '__main__':
    build()
